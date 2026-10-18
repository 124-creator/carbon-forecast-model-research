"""Extract development-only EEX T3PA events, retaining cancellations and missing prices.

This is a source parser, not a forecast dataset selector. It never defines a
next event, never opens 2024/2025 files, and never filters events by success.
Only original annual 2012--2023 workbooks and an approved development daily
CSV are accepted. Decimal strings preserve decoded Excel cached values without
price rounding. Raw typed source rows are written only under local_only/.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import datetime as dt
from decimal import Decimal, InvalidOperation, localcontext
import hashlib
import io
import json
from pathlib import Path
import re
import warnings

YEARS = tuple(range(2012, 2024))
SHEET = "Primary Market Auction"
CONTRACT = "T3PA"
KNOWN_OTHER_CONTRACTS = {"T2PA", "EAA2", "EAA3"}
FAILURE = re.compile(r"cancel|unsuccess|not successful|not held|failed|failure", re.I)
ALIASES = {
    "date": {"date"},
    "auction_time": {"time", "auction time"},
    "auction_name": {"auction name"},
    "contract": {"contract"},
    "certificate_raw": {"certificate"},
    "source_status": {"status"},
    "country_raw": {"country"},
    "auction_details_raw": {"auction details"},
    "official_auction_id_raw": {"auction id", "auction identifier"},
    "clearing_price_eur_per_tco2": {"auction price €/tco2", "auction price eur/tco2"},
    "supply_tco2": {"auction volume tco2", "auction volume"},
    "total_bid_volume_tco2": {"total amount of bids"},
    "cover_ratio_reported": {"cover ratio"},
    "bid_min_eur_per_tco2": {"minimal price €/tco2", "minimum bid €/tco2", "minimum bid eur/tco2"},
    "bid_max_eur_per_tco2": {"maximal price €/tco2", "maximum bid €/tco2", "maximum bid eur/tco2"},
    "bid_mean_eur_per_tco2": {"mean €/tco2", "mean price eur/tco2"},
    "bid_median_eur_per_tco2": {"median €/tco2", "median price eur/tco2"},
    "bid_count": {"number of bids submitted"},
    "bidder_count": {"total number of bidders"},
    "successful_bid_count": {"number of successful bids"},
    "successful_bidder_count": {"number of successful bidders", "number of succesful bidders"},
    "average_bids_per_bidder": {"average number of bids per bidder"},
    "average_bid_size_tco2": {"average bid size"},
    "average_bid_volume_per_bidder_tco2": {"average volume bid per bidder"},
    "bid_volume_std_per_bidder_tco2": {"standard deviation of bid volume per bidder"},
    "average_volume_won_per_bidder_tco2": {"average volume won per bidder"},
    "volume_won_std_per_bidder_tco2": {"standard deviation of volume won per bidder"},
    "revenue_eur": {"total revenue €", "total revenue eur"},
}
NUMERIC = tuple(k for k in ALIASES if k not in {
    "date", "auction_time", "auction_name", "contract", "certificate_raw",
    "source_status", "country_raw", "auction_details_raw", "official_auction_id_raw",
})
REQUIRED_HEADERS = {
    "date", "auction_time", "auction_name", "contract",
    "clearing_price_eur_per_tco2", "supply_tco2", "total_bid_volume_tco2",
    "cover_ratio_reported", "bid_min_eur_per_tco2", "bid_max_eur_per_tco2",
    "bid_mean_eur_per_tco2", "bid_median_eur_per_tco2", "bidder_count",
    "successful_bidder_count", "revenue_eur",
}
FIELDS = [
    "event_id", "date", "auction_time", "platform", "auction_name", "contract",
    "certificate_raw", "official_auction_id_raw", "country_raw", "source_status",
    "event_status", "status_column_present", "price_observed",
    *NUMERIC, "auction_details_raw", "failure_markers_json", "missing_fields",
    "absent_source_columns", "nonnumeric_fields", "anomaly_flags",
    "completed_match_eligible", "old_completion_evidence", "old_component_id",
    "duplicate_event_identity", "within_date_event_count", "same_time_event_count",
    "source_year", "source_file", "source_sheet", "source_excel_row",
    "source_header_excel_row", "source_file_sha256", "source_row_sha256",
]


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def norm(value):
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip().lower()


def dump(obj):
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def typed(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, float) and not (float("-inf") < value < float("inf")):
        return {"type": "nonfinite_excel_numeric", "value": str(value)}
    return value


def text_cell(value):
    if value is None:
        return ""
    if isinstance(value, dict):
        return dump(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def decimal_value(value):
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        return None
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def decimal_text(value):
    if value is None:
        return ""
    value = value if isinstance(value, Decimal) else decimal_value(value)
    if value is None:
        return ""
    result = format(value, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return result or "0"


def date_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return dt.datetime.strptime(value.strip(), pattern).date().isoformat()
            except ValueError:
                continue
    raise ValueError("Unrecognized source date format")


def time_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, dt.datetime):
        return value.time().isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if isinstance(value, str):
        for pattern in ("%H:%M:%S", "%H:%M"):
            try:
                return dt.datetime.strptime(value.strip(), pattern).time().isoformat()
            except ValueError:
                continue
    raise ValueError("Unrecognized source time format; no timezone is inferred")


def platform_from_name(name):
    value = re.sub(r"[-()]", " ", name.upper())
    tokens = value.split()
    if "NIR" in tokens or "NI" in tokens:
        return "NI"
    if "DE" in tokens:
        return "DE"
    if "PL" in tokens:
        return "PL"
    if "EU" in tokens or "TCAP" in tokens:
        return "CAP"
    return "UNKNOWN"


def read_source(path):
    """Decode only the primary worksheet of one explicitly allowed annual file."""
    match = re.fullmatch(r"emission-spot-primary-market-auction-report-(20\d\d)-data\.(xls|xlsx)", path.name)
    if not match or int(match.group(1)) not in YEARS:
        raise ValueError("Only explicitly named 2012--2023 originals may be read")
    blob = path.read_bytes()
    if path.suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=blob, on_demand=True)
        sheet = book.sheet_by_name(SHEET)
        rows = []
        for r in range(sheet.nrows):
            cells = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = {"excel_error_code": int(cell.value)}
                else:
                    value = cell.value
                cells.append(value)
            rows.append(cells)
        book.release_resources()
    else:
        import openpyxl
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style", category=UserWarning)
            book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            sheet = book[SHEET]
            rows = []
            for row in sheet.iter_rows():
                rows.append([{"excel_error": c.value} if c.data_type == "e" else c.value for c in row])
            book.close()
    return blob, rows


def header_map(rows):
    found = []
    for index, row in enumerate(rows[:12]):
        values = [norm(c) for c in row]
        if "date" in values and "contract" in values:
            found.append((index, values))
    if len(found) != 1:
        raise ValueError("Ambiguous primary worksheet header")
    index, values = found[0]
    columns = {}
    for field, aliases in ALIASES.items():
        hits = [i for i, value in enumerate(values) if value in aliases]
        if len(hits) > 1 or (not hits and field in REQUIRED_HEADERS):
            raise ValueError(f"Absent or ambiguous header for {field}")
        columns[field] = hits[0] if hits else None
    return index, columns


def get_cell(row, column):
    return None if column is None or column >= len(row) else row[column]


def classify(record, source_row, status_present):
    values = {field: decimal_value(record[field]) for field in NUMERIC}
    markers = [text_cell(value) for value in source_row if isinstance(value, str) and FAILURE.search(value)]
    flags = []
    p, supply, revenue = (values[k] for k in ("clearing_price_eur_per_tco2", "supply_tco2", "revenue_eur"))
    complete_positive = all(v is not None and v > 0 for v in (p, supply, revenue))
    with localcontext() as context:
        context.prec = 50
        revenue_ok = bool(complete_positive and abs(revenue - p * supply) <= Decimal("0.01"))
    status = norm(record["source_status"])
    if status_present:
        if status == "successful":
            event_status = "reported_successful"
        elif FAILURE.search(status):
            event_status = "reported_cancelled_or_failed"
        else:
            event_status = "reported_status_other_or_missing"
            flags.append("unrecognized_or_missing_reported_status")
    elif markers:
        event_status = "legacy_failure_marker"
    elif p is not None:
        event_status = "legacy_price_observed_status_absent"
    else:
        event_status = "legacy_missing_price_status_absent"
    if not record["date"]:
        flags.append("date_missing")
    if not record["auction_time"]:
        flags.append("time_missing")
    if record["platform"] == "UNKNOWN":
        flags.append("unmapped_platform")
    if record["certificate_raw"] and record["certificate_raw"].strip() != "EUA":
        flags.append("certificate_not_eua")
    if p is None:
        flags.append("clearing_price_missing_or_nonnumeric")
    if complete_positive and not revenue_ok:
        flags.append("revenue_identity_mismatch_over_one_cent")
    for field, value in values.items():
        if value is not None and value < 0:
            flags.append(f"negative:{field}")
    for field in ("bid_count", "bidder_count", "successful_bid_count", "successful_bidder_count"):
        value = values[field]
        if value is not None and value != value.to_integral_value():
            flags.append(f"noninteger:{field}")
    for small, large in (("successful_bid_count", "bid_count"), ("successful_bidder_count", "bidder_count")):
        a, b = values[small], values[large]
        if a is not None and b is not None and a > b:
            flags.append(f"exceeds_total:{small}")
    lo, hi = values["bid_min_eur_per_tco2"], values["bid_max_eur_per_tco2"]
    if lo is not None and hi is not None and lo > hi:
        flags.append("bid_min_exceeds_bid_max")
    eligible = bool(complete_positive and revenue_ok and not markers and
                    (status == "successful" if status_present else True) and
                    "certificate_not_eua" not in flags and record["date"])
    record.update(event_status=event_status, status_column_present=str(status_present).lower(),
                  price_observed=str(p is not None).lower(), failure_markers_json=dump(markers),
                  anomaly_flags="|".join(flags), completed_match_eligible=str(eligible).lower(),
                  old_completion_evidence=("explicit_successful" if status_present else "inferred_completed") if eligible else "")


def parse_sources(source_dir):
    records, raw_records, inventory = [], [], []
    for year in YEARS:
        suffix = "xls" if year < 2020 else "xlsx"
        path = source_dir / f"emission-spot-primary-market-auction-report-{year}-data.{suffix}"
        blob, rows = read_source(path)
        hi, columns = header_map(rows)
        contract_counts, nonempty_undated, included = Counter(), [], []
        for ri, source_row in enumerate(rows[hi + 1:], start=hi + 2):
            contract = text_cell(get_cell(source_row, columns["contract"])).strip()
            if not contract:
                if any(v not in (None, "") for v in source_row):
                    nonempty_undated.append(ri)
                continue
            if contract not in {CONTRACT, *KNOWN_OTHER_CONTRACTS}:
                raise ValueError(f"Unknown contract in {path.name}, source row {ri}")
            contract_counts[contract] += 1
            if contract != CONTRACT:
                continue
            selected = {key: get_cell(source_row, column) for key, column in columns.items()}
            date = date_text(selected["date"])
            if date and (date[:4] != str(year) or not ("2012-01-01" <= date < "2024-01-01")):
                raise ValueError("Date outside the development year: no later-period output is allowed")
            time = time_text(selected["auction_time"])
            name = text_cell(selected["auction_name"]).strip()
            identity = (date, time, name, contract)
            identity_blob = json.dumps(identity, ensure_ascii=False, separators=(",", ":")).encode()
            raw_payload = {"source_file": path.name, "source_sheet": SHEET, "source_excel_row": ri,
                           "cells": [typed(v) for v in source_row]}
            row_hash = sha256(dump(raw_payload).encode())
            record = {key: decimal_text(selected[key]) if key in NUMERIC else text_cell(selected[key])
                      for key in ALIASES}
            record.update(event_id="eex_event_" + sha256(identity_blob)[:24], date=date, auction_time=time,
                          platform=platform_from_name(name), auction_name=name, contract=contract,
                          source_year=str(year), source_file=path.name, source_sheet=SHEET,
                          source_excel_row=str(ri), source_header_excel_row=str(hi + 1),
                          source_file_sha256=sha256(blob), source_row_sha256=row_hash,
                          old_component_id="eex_" + sha256(identity_blob)[:24])
            absent = [key for key, col in columns.items() if col is None]
            missing = [key for key in NUMERIC if decimal_value(selected[key]) is None]
            nonnumeric = [key for key in missing if columns[key] is not None and selected[key] not in (None, "")]
            record.update(absent_source_columns="|".join(absent), missing_fields="|".join(missing),
                          nonnumeric_fields="|".join(nonnumeric))
            classify(record, source_row, columns["source_status"] is not None)
            raw_records.append({"event_id": record["event_id"], "source_row_sha256": row_hash, **raw_payload})
            records.append(record)
            included.append(record)
        inventory.append({"year": year, "file": path.name, "bytes": len(blob), "sha256": sha256(blob),
                          "sheet": SHEET, "header_excel_row": hi + 1,
                          "header": [text_cell(v) for v in rows[hi]],
                          "field_columns_one_based": {key: value + 1 if value is not None else None for key, value in columns.items()},
                          "all_contract_rows": dict(contract_counts), "included_t3pa_rows": len(included),
                          "platform_counts": dict(Counter(x["platform"] for x in included)),
                          "event_status_counts": dict(Counter(x["event_status"] for x in included)),
                          "source_status_counts": dict(Counter(x["source_status"] or "<empty_or_absent>" for x in included)),
                          "missing_numeric_field_counts": {key: sum(not r[key] for r in included) for key in NUMERIC},
                          "nonempty_rows_without_contract": nonempty_undated})
    identities = Counter(r["event_id"] for r in records)
    dates = Counter(r["date"] for r in records)
    times = Counter((r["date"], r["auction_time"]) for r in records)
    for record in records:
        record.update(duplicate_event_identity=str(identities[record["event_id"]] > 1).lower(),
                      within_date_event_count=str(dates[record["date"]]),
                      same_time_event_count=str(times[(record["date"], record["auction_time"])]) )
    records.sort(key=lambda r: (r["date"] or "9999", r["auction_time"], r["source_file"], int(r["source_excel_row"])))
    return records, raw_records, inventory


def verify_source_identity(source_dir, csv_path, inventory):
    """One full source-row reread after CSV serialization; no training or filters."""
    with csv_path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["source_file"]].append(row)
    indexed = {x["file"]: x for x in inventory}
    comparisons, failures, rows_checked = 0, [], 0
    for filename, extracted in grouped.items():
        blob, original = read_source(source_dir / filename)
        meta = indexed[filename]
        if sha256(blob) != meta["sha256"]:
            raise ValueError("Source file changed between parse and identity review")
        for row in extracted:
            ri = int(row["source_excel_row"])
            cells = original[ri - 1]
            payload = {"source_file": filename, "source_sheet": SHEET, "source_excel_row": ri,
                       "cells": [typed(v) for v in cells]}
            if sha256(dump(payload).encode()) != row["source_row_sha256"]:
                failures.append({"file": filename, "row": ri, "field": "source_row_sha256"})
            comparisons += 1
            for field, column in meta["field_columns_one_based"].items():
                cell = cells[column - 1] if column is not None else None
                if field == "date":
                    expected = date_text(cell)
                elif field == "auction_time":
                    expected = time_text(cell)
                elif field in NUMERIC:
                    expected = decimal_text(cell)
                else:
                    expected = text_cell(cell).strip() if field in {"auction_name", "contract"} else text_cell(cell)
                comparisons += 1
                if row[field] != expected:
                    failures.append({"file": filename, "row": ri, "field": field})
            rows_checked += 1
    return {"status": "PASS" if not failures else "FAIL", "scope": "All extracted source rows and every mapped source field, after CSV round trip; same Excel reader libraries, source hashes and declared header mappings.",
            "rows_checked": rows_checked, "cell_or_identity_comparisons": comparisons,
            "failures": failures, "source_files_reread": len(grouped)}


def verify_old_daily(records, old_daily_path):
    """Compare only completed, exactly identifiable component sets with old daily."""
    with old_daily_path.open(encoding="utf-8", newline="") as stream:
        old_rows = list(csv.DictReader(stream))
    if any(not ("2012-01-01" <= x["date"] < "2024-01-01") for x in old_rows):
        raise ValueError("Old daily comparison may use development rows only")
    if len({x["date"] for x in old_rows}) != len(old_rows):
        raise ValueError("Duplicate dates in old daily comparison")
    grouped = defaultdict(list)
    for row in records:
        if row["completed_match_eligible"] == "true":
            grouped[row["date"]].append(row)
    checked, identities_checked, failures, unmatched = 0, 0, [], []
    max_error = defaultdict(lambda: Decimal(0))
    compared_days = set()
    with localcontext() as context:
        context.prec = 50
        for old in old_rows:
            group = grouped.get(old["date"], [])
            actual_ids = sorted(r["old_component_id"] for r in group)
            expected_ids = sorted(old["component_ids"].split("|"))
            if not group or actual_ids != expected_ids or len(set(actual_ids)) != len(actual_ids):
                unmatched.append({"date": old["date"], "reason": "component_identity_set_not_matched"})
                continue
            identities_checked += len(group)
            compared_days.add(old["date"])
            supply = sum(Decimal(x["supply_tco2"]) for x in group)
            def weighted(field):
                if any(not x[field] for x in group):
                    return None
                return sum(Decimal(x[field]) * Decimal(x["supply_tco2"]) for x in group) / supply
            total_bids = sum(Decimal(x["total_bid_volume_tco2"]) for x in group) if all(x["total_bid_volume_tco2"] for x in group) else None
            computed = {"price": weighted("clearing_price_eur_per_tco2"), "auction_volume": supply,
                        "total_bid_volume": total_bids, "cover_ratio": total_bids / supply if total_bids is not None else None,
                        "bidder_count": weighted("bidder_count"), "successful_bidder_count": weighted("successful_bidder_count"),
                        "auction_count": Decimal(len(group))}
            for field, value in computed.items():
                expected = decimal_value(old[field])
                if (expected is None) != (value is None):
                    failures.append({"date": old["date"], "field": field, "reason": "missingness_differs"})
                elif expected is not None:
                    error = abs(value - expected)
                    max_error[field] = max(max_error[field], error)
                    if error > Decimal("1e-40"):
                        failures.append({"date": old["date"], "field": field, "reason": "numeric_mismatch"})
                checked += 1
    return {"status": "PASS" if not failures and not unmatched else "PARTIAL_OR_FAIL",
            "old_daily_file_sha256": sha256(old_daily_path.read_bytes()), "old_daily_rows": len(old_rows),
            "completed_events_matched_by_component_id": identities_checked, "matched_daily_rows": len(compared_days),
            "multi_event_days_matched": sum(len(grouped[d]) > 1 for d in compared_days),
            "numeric_or_missingness_comparisons": checked, "max_absolute_error_by_field": {k: str(v) for k, v in max_error.items()},
            "absolute_tolerance": "1e-40", "unmatched_old_daily_rows": unmatched,
            "eligible_event_dates_not_in_old_daily": sorted(set(grouped) - compared_days), "failures": failures,
            "scope": "Comparison-only successful/inferred-completed subset with exact old component IDs. Cancellations remain in the new event CSV and never define next-event eligibility."}


def write_schema(path, inventory):
    numeric_descriptions = {
        "clearing_price_eur_per_tco2": "原表 Auction Price；EUR/tCO2，取消或缺失保留空。",
        "supply_tco2": "原表 Auction Volume，即记号 S；tCO2。实际供给/成交量的事件含义按原表与状态保留，不能自动当作事前公告量。",
        "total_bid_volume_tco2": "原表 Total Amount of Bids，即记号 B；总投标量，非投标笔数。",
        "cover_ratio_reported": "原表 Cover Ratio，保留原报告精度，不用 B/S 覆写。",
        "bid_min_eur_per_tco2": "原表 Minimum Bid/Minimal Price；EUR/tCO2。",
        "bid_max_eur_per_tco2": "原表 Maximum Bid/Maximal Price；EUR/tCO2。",
        "bid_mean_eur_per_tco2": "原表 Mean；EUR/tCO2。加权方式未认证，不认定数量加权。",
        "bid_median_eur_per_tco2": "原表 Median；EUR/tCO2。加权方式未认证，不认定数量加权。",
        "bid_count": "Number of bids submitted；投标笔数；2012–2016 原列缺失。",
        "bidder_count": "Total Number of Bidders；参与者人数，不等于投标笔数。",
        "successful_bid_count": "Number of successful bids；成功投标笔数；2012–2016 原列缺失。",
        "successful_bidder_count": "Number of Successful Bidders；成功参与者人数。",
        "average_bids_per_bidder": "Average number of bids per bidder；2017 起提供。",
        "average_bid_size_tco2": "Average bid size；投标数量，tCO2；2017 起提供。",
        "average_bid_volume_per_bidder_tco2": "Average volume bid per bidder；tCO2；2017 起提供。",
        "bid_volume_std_per_bidder_tco2": "Standard deviation of bid volume per bidder；数量标准差，tCO2，绝不是价格标准差。",
        "average_volume_won_per_bidder_tco2": "Average volume won per bidder；tCO2；2017 起提供。",
        "volume_won_std_per_bidder_tco2": "Standard deviation of volume won per bidder；数量标准差，tCO2，绝不是价格标准差。",
        "revenue_eur": "原表 Total Revenue；EUR，不用价格乘量覆写。",
    }
    lines = ["# EEX T3PA 事件级开发数据字段", "", "本文件由 `parse_auction_events.py` 生成。只读 2012–2023 的 12 个年度原件；每一行是原表一条 T3PA 拍卖记录。保留取消、缺失价格和所有原表异常，不按成功状态筛选事件，不生成下一事件标签。", "",
             "## 来源、平台与时间", "", "来源为既有已批准 EEX 历史拍卖档案的开发年度成员；精确文件名、字节哈希、表头、列索引见 `DATA_QUALITY.json`。原件可定位于官方档案：https://www.eex.com/fileadmin/EEX/Downloads/Markets/Environmentals/EUA_Emission_Spot_Primary_Market_Auction_Report/Archive_Reports/emission-spot-primary-market-auction-report-2012-2025-data.zip 。本程序不下载或打开完整跨确认年份档案。", "",
             "`platform` 统一为 CAP、DE、PL、NI：EU/TCAP/CAP2-EU/CAP3 EU → CAP；德国 → DE；波兰及 CAP2-PL/CAP3 PL → PL；CAP3 NIR → NI。`auction_name` 与 `country_raw` 保留原文，平台制度代际未被删去。", "",
             "`date`、`auction_time` 分别来自原表日期与 Time/Auction Time。Excel 时钟原样解码；不据此认证 UTC、首次结果发布时刻或事前公告。相同日期多拍卖保留独立行；同一时钟值不虚构可观察先后。`same_time_event_count` 提醒需要同时事件处理。排序只是文件布局，不代表信息先后已认证。", "",
             "## 数值字段", "", "空字符串表示源列不存在、原单元为空或不是有限数值；由 absent_source_columns、missing_fields、nonnumeric_fields 区分，原始单元保存在私有 JSONL。数值由 Excel 缓存值解码为十进制文本，不插补、不平滑、不舍入价格、不更正异常。", "", "| 字段 | 含义 |", "|---|---|"]
    lines.extend(f"| `{key}` | {numeric_descriptions[key]} |" for key in NUMERIC)
    lines.extend(["", "## 状态与标识", "",
                  "- `event_id` 是日期、时钟、Auction Name、Contract 的哈希标识，不是假定存在的官方 ID；`official_auction_id_raw` 在本批原表中缺列时为空。",
                  "- `source_status` 保留原文；2012–2019 没有 Status 列。`event_status` 区分 reported_successful、reported_cancelled_or_failed、legacy_price_observed_status_absent 等，不把早期推断说成官方成功标签。",
                  "- `status_column_present`、`price_observed` 是原表事后观测标志；`failure_markers_json`、`anomaly_flags` 保留取消/异常标记。取消场次价格空值不能填零。",
                  "- `completed_match_eligible`、`old_completion_evidence`、`old_component_id` 只用于与旧日表核对；不得用这些事后标志定义下一次可预测拍卖或筛选预报签发事件。",
                  "- `duplicate_event_identity`、`within_date_event_count`、`same_time_event_count` 是结构标记。发现重复不静默删除；来源文件、工作表、Excel 行号始终保留。",
                  "- `source_year/source_file/source_sheet/source_excel_row/source_header_excel_row/source_file_sha256/source_row_sha256` 给出精确定位。私有 `auction_event_source_rows.jsonl` 保存该来源行全部解码单元，包括取消说明、源国别分项、原文本与 Excel 错误。",
                  "", "## 信息边界与公开范围", "",
                  "本行清算价、实际投标量、覆盖率、投标价格统计、人数、笔数、实际结果状态与收入是本次拍卖结果信息，不能直接作为本次拍卖开始前的输入。事前供给须由独立日历版本按发布时点连接；这里不会从实际成功/取消结果回填已公告日历。研究模型需自行按已冻结截止规则构造过去已公布输入。",
                  "", "本解析器不把 bid-volume 标准差当价格标准差，不认证 Mean/Median 的加权方式，也不把未来成功标签作为下一事件的定义。没有价格分布统计、模型拟合或确认期读取。",
                  "", "公开交付：本脚本、DATA_SCHEMA.md、DATA_QUALITY.json。私有：local_only/ 下事件 CSV 与原始行 JSONL。质量收据只输出源身份、日期/计数、空值与状态及核对误差，不公开逐次价格/投标数值。",
                  "", "## 年份与行数", "", "| 年份 | T3PA 事件行数（保留取消） | 平台 |", "|---|---:|---|"])
    for row in inventory:
        lines.append(f"| {row['year']} | {row['included_t3pa_rows']} | {json.dumps(row['platform_counts'], ensure_ascii=False)} |")
    lines.extend(["", "## 运行", "", "```bash", "python parse_auction_events.py --source-dir /absolute/path/eex_development_originals --output-dir /absolute/path/new_study --old-daily /absolute/path/EEX_EUA_T3PA_daily_development.csv", "```", "", "输出目录中允许存在其他研究文件；只写本解析器负责的四个交付文件及 local_only 的忽略规则。重写已有解析输出需显式 `--overwrite-owned-outputs`。"])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--old-daily", type=Path, required=True)
    parser.add_argument("--overwrite-owned-outputs", action="store_true")
    args = parser.parse_args()
    source_dir, output_dir, old_daily = (p.expanduser().resolve() for p in (args.source_dir, args.output_dir, args.old_daily))
    private = output_dir / "local_only"
    targets = [private / "auction_events_development.csv", private / "auction_event_source_rows.jsonl",
               output_dir / "DATA_SCHEMA.md", output_dir / "DATA_QUALITY.json"]
    if not args.overwrite_owned_outputs and any(p.exists() for p in targets):
        raise FileExistsError("Owned output exists; explicit overwrite is required")
    records, raw_records, inventory = parse_sources(source_dir)
    private.mkdir(parents=True, exist_ok=True)
    ignore = private / ".gitignore"
    if not ignore.exists():
        ignore.write_text("*\n!.gitignore\n", encoding="utf-8")
    with targets[0].open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)
    with targets[1].open("w", encoding="utf-8") as stream:
        for row in raw_records:
            stream.write(dump(row) + "\n")
    write_schema(targets[2], inventory)
    source_review = verify_source_identity(source_dir, targets[0], inventory)
    daily_review = verify_old_daily(records, old_daily)
    flag_counts = Counter(flag for row in records for flag in row["anomaly_flags"].split("|") if flag)
    result = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "script_sha256": sha256(Path(__file__).read_bytes()),
        "scope": {"years": list(YEARS), "contract": CONTRACT, "sheet": SHEET,
                  "training_runs": 0, "confirmation_files_opened": 0, "success_filtered_event_csv": False,
                  "future_success_defines_next_event": False, "prices_changed_or_imputed": False,
                  "mean_median_weighting_certified": False, "price_standard_deviation_available": False,
                  "source_epoch_note": "Historical annual snapshot, not authenticated point-in-time results archive."},
        "output_fields": FIELDS, "source_inventory": inventory,
        "event_rows": len(records), "date_min": min(r["date"] for r in records if r["date"]),
        "date_max": max(r["date"] for r in records if r["date"]),
        "platform_counts": dict(Counter(r["platform"] for r in records)),
        "event_status_counts": dict(Counter(r["event_status"] for r in records)),
        "missing_numeric_field_counts": {key: sum(not r[key] for r in records) for key in NUMERIC},
        "anomaly_flag_counts": dict(flag_counts),
        "anomaly_identity_only": [{k: r[k] for k in ("event_id", "date", "platform", "source_file", "source_excel_row", "event_status", "anomaly_flags")}
                                  for r in records if r["anomaly_flags"]],
        "duplicate_identity_rows": sum(r["duplicate_event_identity"] == "true" for r in records),
        "multi_event_dates": len({r["date"] for r in records if int(r["within_date_event_count"]) > 1}),
        "same_time_multiple_event_groups": len({(r["date"], r["auction_time"]) for r in records if int(r["same_time_event_count"]) > 1}),
        "source_row_identity_review": source_review, "old_daily_aggregation_review": daily_review,
        "private_outputs": [{"relative_path": p.relative_to(output_dir).as_posix(), "bytes": p.stat().st_size, "sha256": sha256(p.read_bytes())} for p in targets[:2]],
        "overall_status": "PASS" if source_review["status"] == "PASS" and daily_review["status"] == "PASS" else "REVIEW_REQUIRED",
        "limitations": ["Cancelled rows and missing prices are retained; historic source may omit other non-held auctions.",
                        "Statistics are as reported; bid-price Mean/Median weighting remains unverified.",
                        "Bid-volume standard deviations are quantities, never price dispersion.",
                        "No release timestamp or future scheduled event was inferred from final success labels.",
                        "Source-row review reuses Excel reader libraries and disclosed header mappings; it is not a second independent parser implementation."]}
    targets[3].write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"overall_status": result["overall_status"], "event_rows": result["event_rows"],
                      "platform_counts": result["platform_counts"], "event_status_counts": result["event_status_counts"],
                      "source_review": {k: source_review[k] for k in ("status", "rows_checked", "cell_or_identity_comparisons")},
                      "daily_review": {k: daily_review[k] for k in ("status", "matched_daily_rows", "completed_events_matched_by_component_id", "multi_event_days_matched")}}, ensure_ascii=False))
    if result["overall_status"] != "PASS":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
