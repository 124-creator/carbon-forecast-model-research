"""Offline byte/row verification of the licensed author-data packet; no training."""
import csv
import hashlib
import io
import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import from_excel


def verify(root):
    root = Path(root).resolve()
    manifest = json.loads((root / 'research/pro_data_packet_20261013/DATA_PACKET_MANIFEST.json').read_text(encoding='utf-8'))
    entries = manifest['files']
    if len(entries) != 6 or len({e['path'] for e in entries}) != 6:
        raise ValueError('Require all six distinct original/CSV inputs')
    checked = {}
    pairs = {}
    for entry in entries:
        path = root / entry['path']
        if path.is_symlink() or not path.resolve().is_relative_to(root):
            raise ValueError('Source escapes packet root')
        body = path.read_bytes()
        if len(body) != entry['bytes'] or hashlib.sha256(body).hexdigest() != entry['sha256']:
            raise ValueError('Byte/hash mismatch: ' + entry['path'])
        checked[entry['path']] = body
        pair = pairs.setdefault(entry['source_filename'], {})
        if entry['kind'] in pair:
            raise ValueError('Duplicate source role')
        pair[entry['kind']] = entry
    if len(pairs) != 3 or any(set(p) != {'original', 'derived_csv'} for p in pairs.values()):
        raise ValueError('Incomplete original/derived source coverage')
    results = []
    for name, pair in pairs.items():
        raw, derived = pair['original'], pair['derived_csv']
        wb = openpyxl.load_workbook(io.BytesIO(checked[raw['path']]), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        if ws.max_column != 2:
            raise ValueError('Unexpected workbook columns')
        values = list(ws.values)
        if values[0] != ('date', 'Y'):
            raise ValueError('Unexpected source header')
        expected = {}
        for rowno, (d, price) in enumerate(values[1:], 2):
            if isinstance(d, datetime):
                d = d.date()
            elif isinstance(d, date):
                pass
            elif isinstance(d, (float, int)):
                d = from_excel(d, wb.epoch).date()
            else:
                d = datetime.fromisoformat(str(d)).date()
            p = Decimal(str(price))
            if not p.is_finite() or p <= 0:
                raise ValueError('Invalid source price')
            expected[rowno] = (d.isoformat(), p)
        wb.close()
        rows = list(csv.DictReader(io.StringIO(checked[derived['path']].decode('utf-8'))))
        if len(rows) != raw['rows'] or len(rows) != derived['rows'] or len(rows) != len(expected):
            raise ValueError('Observation count mismatch')
        seen = set()
        for row in rows:
            rowno = int(row['source_excel_row'])
            if rowno in seen or (row['date'], Decimal(row['author_reported_price'])) != expected.get(rowno):
                raise ValueError('Derived price/date differs from source')
            seen.add(rowno)
        dates = [r['date'] for r in rows]
        if dates != sorted(dates) or len(set(dates)) != len(dates):
            raise ValueError('Dates not strictly ascending')
        results.append({'source_file': name, 'rows': len(rows), 'date_min': dates[0], 'date_max': dates[-1], 'source_date_price_match': True})
    return {'verified_files': len(entries), 'source_workbooks': 3, 'csv_files': 3, 'all_sources': results, 'models_trained_or_scored': False}


if __name__ == '__main__':
    print(json.dumps(verify(Path(__file__).resolve().parents[2]), ensure_ascii=False, indent=2))
