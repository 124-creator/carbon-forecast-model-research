"""Read-only EEX source audit; derive event/day CSVs without modeling."""
from pathlib import Path
from datetime import datetime,timezone
import sys,json,hashlib,importlib.util,warnings,re
import numpy as np
import pandas as pd
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
R=Path(__file__).resolve().parent
ETF=Path(r'E:\Agent员工\碳价格预测研究\独立数据归档\_legacy_202605_archive\主要实验推进_审计加固版_CLEAN_20260528\08-新实验\22_structural_ohlc_q2_route_20260609\03_raw_data\normalized_yahoo_ohlc')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def js(p,obj): p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str,allow_nan=False),encoding='utf-8')
def clean(v):
 if v is None or isinstance(v,str) and not v.strip(): return None
 if isinstance(v,(datetime,pd.Timestamp)): return v.isoformat()
 return v
mapping={'Date':'date','Time':'time','Auction Name':'auction_name','Contract':'contract','Status':'status','Auction Price €/tCO2':'auction_price','Auction Volume tCO2':'auction_volume','Total Amount of Bids':'bid_volume','Cover Ratio':'cover_ratio','Total Number of Bidders':'bidders','Number of Successful Bidders':'successful_bidders','Minimum Bid €/tCO2':'minimum_bid','Maximum Bid €/tCO2':'maximum_bid','Mean €/tCO2':'mean_bid','Median €/tCO2':'median_bid','Country':'country','Total Revenue €':'revenue'}
numcols=[v for k,v in mapping.items() if v not in ['date','time','auction_name','contract','status','country']]
files=[]; rows=[]; anomalies=[]; structures={}; ignored=[]
zipfile=R/'raw/eua_auction_2012_2025.zip'
assert sha(zipfile)=='c3ff6ffaaf348cbd17f61d4726fdbd45a7ead089a0006568a490dd63bb080231'
for p in sorted((R/'raw/extracted').glob('*'))+sorted((R/'raw').glob('*.xlsx')):
 if not p.is_file(): continue
 meta={'path':str(p.relative_to(R)),'sha256':sha(p),'bytes':p.stat().st_size}
 if p.suffix.lower()=='.xls':
  meta.update(parsed=False,engine=None,reason='Bundled Python xlrd absent; legacy workbook not read or converted');files.append(meta);continue
 if p.suffix.lower()!='.xlsx': continue
 with warnings.catch_warnings(record=True) as warns:
  w=openpyxl.load_workbook(p,data_only=True,read_only=True)
  assert len(w.sheetnames)==1
  sheet=w.sheetnames[0]
  vals=list(w[sheet].iter_rows(values_only=True));w.close()
 fw=openpyxl.load_workbook(p,data_only=False,read_only=True)
 formula_count=sum(c.data_type=='f' for row in fw[sheet].iter_rows() for c in row);fw.close()
 hi=next(i for i,row in enumerate(vals[:15]) if 'Date' in row and 'Auction Name' in row)
 headers=[str(z).strip() if z is not None else None for z in vals[hi]]
 localmap=dict(mapping)
 if 'Zone' in headers and 'Country' not in headers: localmap['Zone']=localmap.pop('Country')
 assert all(k in headers for k in localmap),(p.name,headers)
 country_start=headers.index('Country' if 'Country' in headers else 'Zone')+1; count=0
 for lineno,row in enumerate(vals[hi+1:],hi+2):
  nonempty=any(clean(v) is not None for v in row)
  if not nonempty: ignored.append({'source':p.name,'row':lineno,'reason':'entirely empty'});continue
  raw={k:clean(v) for k,v in zip(headers,row) if k}
  rec={localmap[k]:raw.get(k) for k in localmap}
  parsed=pd.to_datetime(rec['date'],errors='coerce')
  if pd.isna(parsed):
   ignored.append({'source':p.name,'row':lineno,'reason':'non-date row retained in ignored_rows receipt','raw':raw});continue
  rec['date']=parsed.strftime('%Y-%m-%d');rec['source_file']=p.name;rec['source_row']=lineno
  rec['country_revenues_json']=json.dumps({k:raw[k] for k in headers[country_start:] if k},ensure_ascii=False)
  rec['raw_fields_json']=json.dumps(raw,ensure_ascii=False)
  for col in numcols:
   original=rec[col]; rec[col]=pd.to_numeric(original,errors='coerce')
   if original is not None and pd.isna(rec[col]): anomalies.append({'source':p.name,'row':lineno,'date':rec['date'],'issue':'nonnumeric_'+col,'value':original})
  rows.append(rec);count+=1
 meta.update(parsed=True,engine='openpyxl '+openpyxl.__version__,sheet=sheet,header_row=hi+1,records=count,formula_cells=formula_count,warnings=[str(z.message) for z in warns]);files.append(meta)
 structures[p.name]={'header_row':hi+1,'headers':headers,'country_revenue_columns':headers[country_start:],'title_nonempty':[[clean(v) for v in rr if clean(v) is not None] for rr in vals[:hi]],'formula_cells':formula_count}
e=pd.DataFrame(rows).sort_values(['date','time','country']).reset_index(drop=True)
e['year']=pd.to_datetime(e.date).dt.year
flags=[]
for i,row in e.iterrows():
 f=[]
 for col in numcols:
  if pd.isna(row[col]):f.append('missing_'+col)
 if pd.isna(row.country): f.append('missing_country')
 if row.cover_ratio<1:f.append('cover_ratio_lt1')
 if row.auction_volume<=0:f.append('auction_volume_nonpositive')
 if row.bid_volume<=0:f.append('bid_volume_nonpositive')
 if row.bidders==0:f.append('zero_bidders')
 if row.successful_bidders==0:f.append('zero_successful_bidders')
 if row.auction_price<=0:f.append('auction_price_nonpositive')
 if row.status!='successful':f.append('status_'+str(row.status))
 if row.minimum_bid>row.maximum_bid:f.append('bid_bounds_reversed')
 if pd.notna(row.auction_price) and pd.notna(row.minimum_bid) and not row.minimum_bid<=row.auction_price<=row.maximum_bid:f.append('clearing_price_outside_bid_range')
 if row.successful_bidders>row.bidders:f.append('successful_bidders_exceed_total')
 if row.auction_volume>0 and pd.notna(row.cover_ratio) and abs(row.cover_ratio-row.bid_volume/row.auction_volume)>.005001:f.append('cover_ratio_formula_mismatch')
 rev=json.loads(row.country_revenues_json); vv=[v for v in rev.values() if isinstance(v,(int,float))]
 if vv and pd.notna(row.revenue) and abs(sum(vv)-row.revenue)>.05:f.append('country_revenue_sum_mismatch')
 if pd.notna(row.revenue) and pd.notna(row.auction_price) and pd.notna(row.auction_volume) and abs(row.revenue-row.auction_price*row.auction_volume)>.05:f.append('revenue_price_volume_mismatch')
 flags.append('|'.join(f))
 for issue in f:anomalies.append({'source':row.source_file,'row':int(row.source_row),'date':row.date,'issue':issue,'value':None})
e['anomaly_flags']=flags
key=['date','time','auction_name','contract','country'];dup=e.duplicated(key,keep=False)
for row in e[dup].itertuples():anomalies.append({'source':row.source_file,'row':row.source_row,'date':row.date,'issue':'duplicate_event_key','value':None})
def strict_sum(s):return float(s.sum()) if s.notna().all() else np.nan
def strict_mean(s):return float(s.mean()) if s.notna().all() else np.nan
def weighted(x,w):
 if x.isna().any() or w.isna().any() or (w<=0).any() or w.sum()<=0:return np.nan
 return float(np.dot(x,w)/w.sum())
def agg(z):
 coverz=z[z.status=='successful']
 a={'n_auctions':len(z),'n_successful':int((z.status=='successful').sum()),'n_non_successful':int((z.status!='successful').sum()),'auction_volume':strict_sum(z.auction_volume),'bid_volume':strict_sum(z.bid_volume),'cover_ratio':weighted(coverz.cover_ratio,coverz.auction_volume),'cover_ratio_weighted':weighted(coverz.cover_ratio,coverz.auction_volume),'cover_ratio_simple_mean':strict_mean(coverz.cover_ratio),'bidders':strict_sum(z.bidders),'successful_bidders':strict_sum(z.successful_bidders),'price':weighted(z.auction_price,z.auction_volume),'price_min':z.minimum_bid.min(skipna=False),'price_max':z.maximum_bid.max(skipna=False),'price_mean':strict_mean(z.mean_bid),'price_median':z.median_bid.median(skipna=False),'bid_dispersion':(z.maximum_bid-z.minimum_bid).mean(skipna=False),'clearing_price_min':z.auction_price.min(skipna=False),'clearing_price_max':z.auction_price.max(skipna=False),'clearing_price_mean':strict_mean(z.auction_price),'clearing_price_median':z.auction_price.median(skipna=False),'n_anomalous_auctions':int(z.anomaly_flags.ne('').sum())}
 a['cover_ratio_from_volumes']=strict_sum(coverz.bid_volume)/strict_sum(coverz.auction_volume) if len(coverz) and strict_sum(coverz.auction_volume)>0 else np.nan
 return a
daily=[]
for date,z in e.groupby('date'):
 a={'date':date,**agg(z),'status_mix':json.dumps(z.status.value_counts().to_dict(),ensure_ascii=False),'country_mix':json.dumps({str(c):{'n_auctions':len(zz),'auction_volume':None if zz.auction_volume.isna().any() else float(zz.auction_volume.sum())} for c,zz in z.groupby('country',dropna=False)},ensure_ascii=False),'auction_records_json':json.dumps(z[['time','auction_name','contract','status','country','source_file','source_row','anomaly_flags']].to_dict('records'),ensure_ascii=False)}
 for country,zz in z.groupby('country',dropna=False):
  for k,v in agg(zz).items():a[str(country)+'__'+k]=v
 for contract,zz in z.groupby('contract'):
  prefix={'T3PA':'EUA','EAA3':'EUAA'}.get(contract,str(contract))
  for k,v in agg(zz).items():a[prefix+'__'+k]=v
  a[prefix+'__country_mix']=json.dumps(zz.country.value_counts().to_dict(),ensure_ascii=False)
 a['contract_mix']=json.dumps(z.contract.value_counts().to_dict(),ensure_ascii=False)
 a['has_EUA']=bool((z.contract=='T3PA').any());a['has_EUAA']=bool((z.contract=='EAA3').any())
 daily.append(a)
panel_all=pd.DataFrame(daily)
panel=panel_all[panel_all.date.between('2020-07-01','2026-06-08')].copy();aligned=[];calmeta=[];unmatched=[]
for asset in ['KRBN','KEUA','KCCA']:
 p=ETF/f'{asset}_ohlc.csv';d=pd.read_csv(p,usecols=['date']); dates=pd.to_datetime(d.date).dt.strftime('%Y-%m-%d');lo,hi=dates.min(),dates.max();observed=set(dates)
 panel[asset+'_observed_session']=panel.date.isin(observed)
 panel[asset+'_within_observed_span']=panel.date.between(lo,hi)
 for dt in panel.loc[~panel.date.isin(observed),'date']:
  unmatched.append({'asset':asset,'auction_date':dt,'has_EUA':bool(panel.loc[panel.date==dt,'has_EUA'].iloc[0]),'has_EUAA':bool(panel.loc[panel.date==dt,'has_EUAA'].iloc[0]),'next_observed_etf_date':min((x for x in observed if x>=dt),default=None),'classification':'before_asset_start' if dt<lo else ('after_asset_end' if dt>hi else 'within_asset_span_not_observed'),'asset_first':lo,'asset_last':hi})
 for y,z in panel.groupby(pd.to_datetime(panel.date).dt.year):
  within=z[z.date.between(lo,hi)];etfy=dates[pd.to_datetime(dates).dt.year==y]
  aligned.append({'asset':asset,'year':int(y),'auction_dates':len(z),'in_etf_calendar':int(z.date.isin(observed).sum()),'not_in_etf_calendar_all':int((~z.date.isin(observed)).sum()),'auction_dates_within_etf_span':len(within),'not_in_etf_calendar_within_span':int((~within.date.isin(observed)).sum()),'etf_observed_dates':len(etfy),'auction_share_of_etf_dates':float(z.date.isin(observed).sum()/len(etfy)) if len(etfy) else None})
 calmeta.append({'asset':asset,'path':str(p),'sha256':sha(p),'first':lo,'last':hi,'rows':len(d),'auction_dates_outside_entire_calendar':int((~panel.date.isin(observed)).sum()),'auction_dates_missing_within_asset_span':int((panel.date.between(lo,hi)&~panel.date.isin(observed)).sum()),'EUA_dates_missing_within_asset_span':int((panel.has_EUA&panel.date.between(lo,hi)&~panel.date.isin(observed)).sum()),'EUA_dates_matched':int((panel.has_EUA&panel.date.isin(observed)).sum())})
annual=e.groupby('year').agg(auctions=('date','size'),dates=('date','nunique'),start=('date','min'),end=('date','max'))
missing={c:int(e[c].isna().sum()) for c in e.columns if c not in ['raw_fields_json','country_revenues_json']}
for name,frame in [('auction_events.csv',e),('daily_auction_panel.csv',panel),('unmatched_auction_dates.csv',pd.DataFrame(unmatched)),('annual_coverage.csv',annual.reset_index()),('etf_calendar_alignment.csv',pd.DataFrame(aligned)),('anomalies.csv',pd.DataFrame(anomalies,columns=['source','row','date','issue','value']))]:frame.to_csv(R/name,index=False,encoding='utf-8-sig',float_format='%.15g')
js(R/'schema_audit.json',structures);js(R/'ignored_rows.json',ignored)
summary={'annual':annual.reset_index().to_dict('records'),'status':e.status.value_counts(dropna=False).to_dict(),'contracts':e.contract.value_counts(dropna=False).to_dict(),'countries':e.country.value_counts(dropna=False).to_dict(),'auction_names':e.auction_name.value_counts(dropna=False).to_dict(),'missing':missing,'event_key_duplicate_rows':int(dup.sum()),'exact_raw_duplicate_rows':int(e.raw_fields_json.duplicated(keep=False).sum()),'multi_auction_dates':panel.loc[panel.n_auctions>1,['date','n_auctions','country_mix']].to_dict('records'),'anomaly_counts':pd.Series([a['issue'] for a in anomalies],dtype=str).value_counts().to_dict(),'etf_calendars':calmeta,'alignment':aligned,'first':e.date.min(),'last':e.date.max(),'events':len(e),'dates':len(panel),'2026_after_etf_cutoff':int(((panel_all.date>'2026-06-08')&(panel_all.date<'2027-01-01')).sum()),'panel_requested_span':['2020-07-01','2026-06-08'],'panel_actual_span':[panel.date.min(),panel.date.max()],'parsed_auction_dates_all':len(panel_all),'time_min':str(e.time.min()),'time_max':str(e.time.max())}
js(R/'audit_metrics.json',summary)
manifest={'created_utc':datetime.now(timezone.utc).isoformat(),'python':sys.executable,'versions':{'pandas':pd.__version__,'openpyxl':openpyxl.__version__},'legacy_xls_engine_available':bool(importlib.util.find_spec('xlrd')),'archive':{'path':str(zipfile.relative_to(R)),'sha256':sha(zipfile),'bytes':zipfile.stat().st_size},'sources':files,'download_2026_url':'https://public.eex-group.com/eex/eua-auction-report/emission-spot-primary-market-auction-report-2026-data.xlsx','source_units':{'price':'EUR per EUA, EUR/tCO2eq','auction_volume':'tCO2eq / allowances','bid_volume':'allowances; validated against bid_volume/auction_volume = cover_ratio within rounding','country_columns':'revenue EUR by recipient; not bid nationality or quantities'},'row_rules':{'date_records_retained':len(e),'daily_rows':len(panel),'non_data_rows':len(ignored),'anomalies_retained':True,'duplicates_retained':True,'missing_filled_with_zero':False,'cancelled_excluded':False,'xls_2012_2019':'unparsed and outside ETF sample scope, not a missing-data claim','panel_requested_span':['2020-07-01','2026-06-08'],'parsed_records_outside_panel_span':int((~e.date.between('2020-07-01','2026-06-08')).sum()),'all_source_date_records_preserved_in':'auction_events.csv','nondata_details':'ignored_rows.json'},'aggregation':{'contract_scope':'Unprefixed aggregates include ALL source contracts; EUA__ prefix is T3PA/general allowances, EUAA__ is EAA3/aviation. Use has_EUA and EUA__ fields for EUA experiments; no event deleted.','auction_volume':'sum offered/source auction volume, NOT certified cleared volume for cancelled events','bid_volume':'sum; missing anywhere makes total missing','cover_ratio':'successful events only; source-volume weighted reported ratio, simple mean and ratio-of-sums; cancelled source ratios retained in auction_events.csv','bidders':'sum event bidder counts, NO cross-auction unique bidder deduplication possible','price':'source auction-volume weighted clearing price; undefined if missing or nonpositive weights','price_min_max':'extremes of submitted bid price ranges across events','price_mean':'unweighted mean of auction-level mean bid prices, NOT exact pooled bid mean','price_median':'median of auction-level bid medians, NOT exact pooled bid median','bid_dispersion':'simple mean of event maximum_bid-minus-minimum_bid ranges, NOT standard deviation','country_mix':'auction mandate Country, counts and volume; EU and country-prefixed aggregate columns preserve separation','absent_country':'blank, not zero','cancelled':'retained; n_successful/n_non_successful and status_mix distinguish outcomes','numeric_missing':'strict complete-event aggregation, no silent partial sums'},'calendar_inputs':calmeta,'outputs':{n:sha(R/n) for n in ['auction_events.csv','daily_auction_panel.csv','unmatched_auction_dates.csv','annual_coverage.csv','etf_calendar_alignment.csv','anomalies.csv','schema_audit.json','ignored_rows.json','audit_metrics.json']},'script_sha256':sha(Path(__file__))}
for m in files:assert sha(R/m['path'])==m['sha256']
js(R/'manifest.json',manifest)
print(json.dumps({k:v for k,v in summary.items() if k not in ['alignment','multi_auction_dates','auction_names','missing']},ensure_ascii=False,indent=2))
