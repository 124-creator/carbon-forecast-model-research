from pathlib import Path
import json, importlib.util, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
R=Path(__file__).resolve().parent
out={}
for p in sorted((R/'raw/extracted').glob('*.xlsx'))+list((R/'raw').glob('*.xlsx')):
 w=openpyxl.load_workbook(p,data_only=True,read_only=True)
 s=w[w.sheetnames[0]]
 out[p.name]={'sheets':w.sheetnames,'rows':s.max_row,'cols':s.max_column,'head':list(s.iter_rows(min_row=1,max_row=9,values_only=True)),'tail':list(s.iter_rows(min_row=max(1,s.max_row-3),max_row=s.max_row,values_only=True))}
(R/'workbook_structure.json').write_text(json.dumps(out,default=str,ensure_ascii=False,indent=2),encoding='utf-8')
for name,z in out.items():
 print(name,z['rows'],z['cols'])
 for j,row in enumerate(z['head']):
  if j in [4,5,6,7,8]: print(j+1,row)
 print('tail',z['tail'])
