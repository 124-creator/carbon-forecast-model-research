"""Preserve author workbook values, split by a pre-acquisition calendar cutoff.

Read-only source extraction: no market fit, returns, price summary, smoothing,
imputation, inverse-log reconstruction, or numerical confirmation output.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import math
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
MANIFEST = BASE / 'INDEPENDENT_DATA_MANIFEST.json'
SOURCES = [
    {
        'series_id': 'SHANGHAI_D4',
        'filename': 'mendeley_d4_carbon.xlsx',
        'sha256': '41592d944375b3109010c151d9b1c103fb906d608afcbfd529d7381f18896697',
        'bytes': 66625,
        'headers': ['TradingDate', 'CityName', '收盘价', '碳配额'],
        'target_column': 2,
        'output_target': 'price',
        'target_definition': 'Author workbook column 收盘价, all CityName values 上海; reported closing-price target, not independently certified official close.',
        'positive_required': True,
        'doi': '10.17632/d4bv52fgmd.1',
        'landing_url': 'https://data.mendeley.com/datasets/d4bv52fgmd/1',
        'download_url': 'https://data.mendeley.com/public-files/datasets/d4bv52fgmd/files/4935504f-cf2d-45c0-add4-d6c20d667a45/file_downloaded',
        'archive_member': 'data/碳价格与配额.xlsx',
        'source_side_processing': 'Not documented by the downloaded README or landing description. No assertion of no historical filling/cleaning. We make no additional price changes.',
        'primary_candidate': True,
    },
    {
        'series_id': 'CEA_LIAO_LOG',
        'filename': 'mendeley_5s_processed.xlsx',
        'sha256': 'f6d9be7c2e4bb4463cd123936882bbca79626312a7f4a718804cd217816b0c06',
        'bytes': 109902,
        'headers': ['交易日期Idxtrd01', 'LNCEA', 'LNVOLUME', 'LNCSI300', 'LNEL300', 'LNEUA', 'LNCOAL', 'LNOIL', 'LNGAS', 'LNELCO2'],
        'target_column': 1,
        'output_target': 'ln_price',
        'target_definition': 'Author-supplied LNCEA numerical values; log-labeled national carbon series. No inverse transform or claim of original unrounded price.',
        'positive_required': False,
        'doi': '10.17632/5skvz8sgrz.1',
        'landing_url': 'https://data.mendeley.com/datasets/5skvz8sgrz/1',
        'download_url': 'https://data.mendeley.com/public-files/datasets/5skvz8sgrz/files/5cb129aa-e29b-4b8f-a3bb-dc3fafb3ffa8/file_downloaded',
        'archive_member': None,
        'source_side_processing': 'Filename Processed data.xlsx and LN-prefixed columns establish transformed author data. Original price precision, vintage/composite identity, and filling protocol remain unverified.',
        'primary_candidate': False,
    },
]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_rows(source: dict) -> tuple[list[tuple[str, int | float, int]], int]:
    path = BASE / 'local_only' / 'raw' / source['filename']
    if path.stat().st_size != source['bytes'] or digest(path) != source['sha256']:
        raise ValueError('Source byte/hash mismatch: ' + source['series_id'])
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if book.sheetnames != ['Sheet1']:
        raise ValueError('Unexpected sheet list')
    sheet = book['Sheet1']
    header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    if header != source['headers']:
        raise ValueError('Unexpected source header')
    rows, blanks = [], 0
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if all(value is None for value in row):
            blanks += 1
            continue
        value = row[0]
        if isinstance(value, dt.datetime):
            date = value.date()
        elif isinstance(value, dt.date):
            date = value
        elif isinstance(value, str):
            date = dt.datetime.strptime(value.strip(), '%Y-%m-%d').date()
        else:
            raise ValueError('Unrecognized source date type')
        if source['series_id'] == 'SHANGHAI_D4' and row[1] != '上海':
            raise ValueError('Unexpected market identity in workbook')
        target = row[source['target_column']]
        if isinstance(target, bool) or not isinstance(target, (int, float)):
            raise ValueError('Target is nonnumeric or a formula')
        if not math.isfinite(target) or (source['positive_required'] and target <= 0):
            raise ValueError('Invalid target; no automatic deletion or imputation')
        rows.append((date.isoformat(), target, number))
    book.close()
    rows.sort(key=lambda item: item[0])
    if len({item[0] for item in rows}) != len(rows):
        raise ValueError('Duplicate source dates; no silent deduplication')
    return rows, blanks


def csv_bytes(rows: list[tuple], target_name: str) -> bytes:
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, lineterminator='\n')
    writer.writerow(['date', target_name, 'source_row'])
    # repr/str preserves the exact numerical value obtained from the workbook.
    for date, target, number in rows:
        writer.writerow([date, repr(target), number])
    return stream.getvalue().encode('utf-8')


def main() -> None:
    if MANIFEST.exists():
        raise FileExistsError('Existing manifest; refusing to rerun or overwrite')
    scope = BASE / 'INDEPENDENT_SOURCE_SCOPE.json'
    if not scope.exists():
        raise RuntimeError('Missing pre-acquisition independent-source scope')
    plans, all_metadata = [], []
    for source in SOURCES:
        rows, blank_rows = source_rows(source)
        groups = {
            'development': [row for row in rows if row[0] <= '2023-12-31'],
            'confirmation_sealed': [row for row in rows if '2024-01-01' <= row[0] <= '2025-12-31'],
        }
        metadata = dict(source)
        metadata.update(license='CC BY 4.0 (depositor declaration)', observed_records=len(rows),
                        date_min=rows[0][0], date_max=rows[-1][0],
                        blank_spreadsheet_rows_skipped=blank_rows,
                        blank_rows_are_not_missing_price_observations=True,
                        duplicate_dates=0, invalid_target_rows=0,
                        experimental_records=sum(len(group) for group in groups.values()),
                        dates_2026_or_later_excluded=sum(row[0] >= '2026-01-01' for row in rows),
                        transformations=['date ISO formatting', 'ascending sort', 'select author target column'],
                        extra_price_transformations=[], price_unit_independently_verified=False,
                        daily_calendar_completeness_verified=False,
                        previous_project_model_use_found=False,
                        previous_project_model_use_audit_scope='Fixed GitHub tree/manifests and targeted prior experiment reports, not a mathematical proof of absence.',
                        partitions={})
        for partition, group in groups.items():
            path = BASE / 'local_only' / partition / (source['series_id'] + '.csv')
            if path.exists():
                raise FileExistsError('Refusing to overwrite ' + path.name)
            body = csv_bytes(group, source['output_target'])
            plans.append((path, body))
            years = {}
            for date, _, _ in group:
                years[date[:4]] = years.get(date[:4], 0) + 1
            metadata['partitions'][partition] = {
                'path': str(path.relative_to(BASE)), 'columns': ['date', source['output_target'], 'source_row'],
                'records': len(group), 'date_min': group[0][0] if group else None,
                'date_max': group[-1][0] if group else None, 'year_counts': years,
                'bytes': len(body), 'sha256': hashlib.sha256(body).hexdigest(),
            }
        all_metadata.append(metadata)
    manifest = {
        'prepared_utc': dt.datetime.now(dt.timezone.utc).isoformat(),
        'scope_path': str(scope.relative_to(BASE)), 'scope_sha256': digest(scope),
        'parser_path': Path(__file__).name, 'parser_sha256': digest(Path(__file__)),
        'series': all_metadata, 'model_fit': False, 'returns_or_price_statistics': False,
        'confirmation_values_printed': False,
        'source_row_is_provenance_only_not_a_feature': True,
        'warning': 'The 760-row LNCEA archive is separate from level-price primary candidates. Do not pool log values into the price field.',
    }
    for path, body in plans:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open('xb') as stream:
            stream.write(body)
    with MANIFEST.open('x', encoding='utf-8') as stream:
        json.dump(manifest, stream, ensure_ascii=False, indent=2)
        stream.write('\n')
    # Metadata contains no price value, price distribution, or forecast score.
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
