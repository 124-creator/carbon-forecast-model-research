"""Install exact study data in a NEW directory, with no model execution.

Public downloads occur only with --download. Local package/cache modes never
fall back to the network. No override for the pinned identities is exposed.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
from pathlib import Path, PurePosixPath
import stat
import urllib.parse
import urllib.request
import zipfile

COMMIT = "1c67e562e721dd82fff4e4143b97a798a1cf44fa"
RAW_BASE = ("https://raw.githubusercontent.com/124-creator/"
            "carbon-forecast-model-research/" + COMMIT +
            "/data/parsed/figshare_28682348/")
RAW_SOURCES = [
    dict(name="GUANGZHOU", filename="guangzhou_reported.csv",
         url=RAW_BASE + "guangzhou_reported.csv", bytes=46549,
         sha256="b75b2c37e158933562d6979ae7532a6323ef9111a9836be8b75e6fe56b3d6573",
         records=1492, start="2018-01-03", end="2024-03-29"),
    dict(name="HUBEI", filename="hubei_reported.csv",
         url=RAW_BASE + "hubei_reported.csv", bytes=45446,
         sha256="1b43a0a30146517ccd4b385777f510a4bcf83395df7ad77c90cdf5476000a71d",
         records=1466, start="2018-01-02", end="2024-03-29"),
    dict(name="EU_REPORTED", filename="eu_reported.csv",
         url=RAW_BASE + "eu_reported.csv", bytes=49203,
         sha256="ff4fa943f8a86407eceb37026bddad2b295f5066c2e66757df6478111fc26b62",
         records=1577, start="2018-01-02", end="2024-03-27"),
    dict(name="SHANGHAI_D4", filename="mendeley_d4_data.zip",
         url="https://data.mendeley.com/public-files/datasets/d4bv52fgmd/files/4935504f-cf2d-45c0-add4-d6c20d667a45/file_downloaded",
         bytes=508904,
         sha256="b4ddba5c3609dfa3c13ed4a0c5998c952225c69dc1bc6c8a8f4c73a5fdfd2862",
         records=1768, start="2013-12-19", end="2025-09-25"),
]
WORKBOOK = dict(member="data/碳价格与配额.xlsx", bytes=66625,
                sha256="41592d944375b3109010c151d9b1c103fb906d608afcbfd529d7381f18896697")
DATA_FILES = [
    dict(name="GUANGZHOU", path="development/GUANGZHOU.csv", bytes=43143,
         sha256="8d78a4f13c5a799519546fd3bfa0e0ac5eaa52b0441a983a95ef06f5e5e34051",
         records=1434, role="development", price_column="price"),
    dict(name="HUBEI", path="development/HUBEI.csv", bytes=42173,
         sha256="04655d59c1b585f6c35c02a43469a965fe3589eb613d7d8d8e46996751b7d797",
         records=1408, role="development", price_column="price"),
    dict(name="EU_REPORTED", path="development/EU_REPORTED.csv", bytes=45699,
         sha256="ca7c192380897076ef57516a11ea2ccefda6ba92ab9823d9aa8f2bb281cefedc",
         records=1515, role="development", price_column="price"),
    dict(name="SHANGHAI_D4", path="development/SHANGHAI_D4.csv", bytes=30443,
         sha256="425424353bb3ac09e3d96b291680b4179e86f180da626a5d19e5c1e58ed960a1",
         records=1507, role="development", price_column="price"),
    dict(name="SHANGHAI_D4", path="confirmation_sealed/SHANGHAI_D4.csv", bytes=5519,
         sha256="5ef3c7664023027f844c6476b40a24c2906d100a35a2415a3d1b4b27ae57ccfc",
         records=261, role="confirmation", price_column="price"),
]


def verify_bytes(blob: bytes, expected: dict, label: str) -> None:
    if len(blob) != expected["bytes"] or hashlib.sha256(blob).hexdigest() != expected["sha256"]:
        raise ValueError("Pinned byte identity mismatch: " + label)


def safe_relative(value: str) -> PurePosixPath:
    p = PurePosixPath(value)
    if (not value or p.is_absolute() or ".." in p.parts or "\\" in value
            or ":" in value or "\x00" in value or value != p.as_posix()):
        raise ValueError("Unsafe relative path")
    return p


def safe_local(root: Path, relative: str) -> Path:
    p = root.joinpath(*safe_relative(relative).parts)
    cursor = root
    for part in safe_relative(relative).parts:
        cursor = cursor / part
        if cursor.is_symlink():
            raise ValueError("Symbolic links are not accepted for source members")
    if not p.is_file() or not p.resolve().is_relative_to(root.resolve()):
        raise FileNotFoundError("Missing regular source file: " + relative)
    return p


def packaged_file(root: Path, spec: dict) -> Path:
    """Accept the canonical layout or the repository's exact flat filenames."""
    alternatives = [spec["path"], spec["name"] + "_" + spec["role"] + ".csv"]
    found = [name for name in alternatives if (root / name).exists() or (root / name).is_symlink()]
    if len(found) != 1:
        raise ValueError("Missing or ambiguous packaged file: " + spec["name"] + "/" + spec["role"])
    return safe_local(root, found[0])


def require_new_output(path: Path) -> Path:
    # Check before resolve so a dangling symlink cannot be used as a new output.
    expanded = path.expanduser()
    if expanded.exists() or expanded.is_symlink():
        raise FileExistsError("Output already exists; use a new directory")
    resolved = expanded.resolve()
    if resolved.exists():
        raise FileExistsError("Output already exists; use a new directory")
    return resolved


def write_new(path: Path, blob: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as stream:
        stream.write(blob)


def json_bytes(value: object) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode()


class FixedPublicRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        original = urllib.parse.urlsplit(req.full_url)
        target = urllib.parse.urlsplit(newurl)
        allowed = {"raw.githubusercontent.com", "data.mendeley.com"}
        if (target.scheme != "https" or target.hostname not in allowed
                or target.hostname != original.hostname or target.username or target.password):
            raise RuntimeError("Unexpected public-source redirect; stopped without fallback")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch_source(source: dict) -> tuple[bytes, dict]:
    opener = urllib.request.build_opener(FixedPublicRedirect())
    request = urllib.request.Request(source["url"], headers={"User-Agent": "CarbonForecastResearch/1.0"})
    with opener.open(request, timeout=40) as response:
        if response.status != 200:
            raise RuntimeError("Unexpected public-source HTTP status")
        blob = response.read(source["bytes"] + 1)
        metadata = dict(status=response.status, final_url=response.url)
    verify_bytes(blob, source, source["filename"])
    return blob, metadata


def reported_development(blob: bytes) -> tuple[bytes, dict]:
    """Keep original cell strings; reproduce the frozen LF DictWriter output."""
    reader = csv.DictReader(io.StringIO(blob.decode("utf-8-sig"), newline=""))
    expected = ["date", "author_reported_price", "source_excel_row", "source_date_kind"]
    if reader.fieldnames != expected:
        raise ValueError("Unexpected reported-CSV header")
    rows = list(reader)
    dates = []
    for row in rows:
        if set(row) != set(expected) or any(value is None for value in row.values()):
            raise ValueError("Malformed reported-CSV row")
        date = dt.date.fromisoformat(row["date"])
        if date.isoformat() != row["date"]:
            raise ValueError("Noncanonical source date")
        value = float(row["author_reported_price"])
        if not math.isfinite(value) or value <= 0:
            raise ValueError("Invalid reported price; no automatic repair")
        dates.append(row["date"])
    if not dates or any(b <= a for a, b in zip(dates, dates[1:])):
        raise ValueError("Reported dates are not unique and ascending")
    selected = [row for row in rows if row["date"] < "2024-01-01"]
    stream = io.StringIO(newline="")
    fields = ["date", "price", "source_excel_row", "source_date_kind"]
    writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    for row in selected:
        writer.writerow(dict(date=row["date"], price=row["author_reported_price"],
                             source_excel_row=row["source_excel_row"], source_date_kind=row["source_date_kind"]))
    return stream.getvalue().encode(), dict(source_records=len(rows), development_records=len(selected),
        source_start=dates[0], source_end=dates[-1],
        excluded_previously_available_suffix_records=len(rows)-len(selected))


def extract_shanghai(archive_blob: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(archive_blob)) as archive:
        members = archive.infolist()
        if len({m.filename for m in members}) != len(members):
            raise ValueError("Duplicate ZIP member")
        if sum(m.file_size for m in members) > 40_000_000:
            raise ValueError("ZIP expansion limit")
        for member in members:
            safe_relative(member.filename.rstrip("/"))
            if stat.S_ISLNK(member.external_attr >> 16):
                raise ValueError("ZIP symbolic link rejected")
        member = archive.getinfo(WORKBOOK["member"])
        if member.file_size != WORKBOOK["bytes"]:
            raise ValueError("Unexpected workbook size")
        with archive.open(member) as stream:
            blob = stream.read(WORKBOOK["bytes"] + 1)
    verify_bytes(blob, WORKBOOK, "Shanghai workbook")
    return blob


def shanghai_partitions(workbook_blob: bytes) -> dict[str, bytes]:
    """Same date/row/repr semantics as the frozen independent-data parser."""
    import openpyxl
    if openpyxl.__version__ != "3.1.5":
        raise RuntimeError("Raw-workbook reproduction requires openpyxl==3.1.5")
    book = openpyxl.load_workbook(io.BytesIO(workbook_blob), read_only=True, data_only=False)
    try:
        if book.sheetnames != ["Sheet1"]:
            raise ValueError("Unexpected Shanghai workbook sheets")
        sheet = book["Sheet1"]
        if list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))) != ["TradingDate", "CityName", "收盘价", "碳配额"]:
            raise ValueError("Unexpected Shanghai workbook header")
        rows = []
        for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if all(value is None for value in row):
                continue
            if row[1] != "上海":
                raise ValueError("Unexpected workbook market")
            date = row[0]
            if isinstance(date, dt.datetime): date = date.date()
            elif isinstance(date, str): date = dt.datetime.strptime(date.strip(), "%Y-%m-%d").date()
            if not isinstance(date, dt.date):
                raise ValueError("Invalid workbook date")
            value = row[2]
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
                raise ValueError("Invalid workbook price; no automatic repair")
            rows.append((date.isoformat(), value, source_row))
    finally:
        book.close()
    rows.sort(key=lambda item: item[0])
    if len({row[0] for row in rows}) != len(rows):
        raise ValueError("Duplicate workbook date; no automatic deduplication")
    groups = {"development": [r for r in rows if r[0] < "2024-01-01"],
              "confirmation_sealed": [r for r in rows if "2024-01-01" <= r[0] < "2026-01-01"]}
    output = {}
    for partition, group in groups.items():
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["date", "price", "source_row"])
        for date, value, row_number in group:
            writer.writerow([date, repr(value), row_number])
        output[partition + "/SHANGHAI_D4.csv"] = stream.getvalue().encode()
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    modes = parser.add_mutually_exclusive_group(required=True)
    modes.add_argument("--package-dir", type=Path, help="Canonical five-file layout or repository data/*_development.csv and *_confirmation.csv")
    modes.add_argument("--source-cache", type=Path, help="Contains the four exact RAW_SOURCES filenames")
    modes.add_argument("--download", action="store_true", help="Fetch only the four pinned public URLs")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    output = require_new_output(args.output)
    mode = "package" if args.package_dir else "raw-cache" if args.source_cache else "public-download"
    if args.dry_run:
        print(json.dumps(dict(output=str(output), mode=mode, creates_files=False, network_calls=0,
                              expected_files=DATA_FILES, public_sources=RAW_SOURCES), ensure_ascii=False, indent=2))
        return
    output.mkdir(parents=True, exist_ok=False)
    write_new(output / ".gitignore", b"local_only/\ndevelopment/\nconfirmation_sealed/\n__pycache__/\n*.pyc\n")
    receipt = dict(started_utc=dt.datetime.now(dt.timezone.utc).isoformat(), mode=mode,
                   files=[], acquisitions=[], source_metadata=[], network_retries=0,
                   model_fit=False, price_statistics_printed=False, confirmation_values_printed=False)
    try:
        derived = {}
        if args.package_dir:
            package = args.package_dir.expanduser().resolve()
            for spec in DATA_FILES:
                blob = packaged_file(package, spec).read_bytes()
                verify_bytes(blob, spec, spec["path"])
                derived[spec["path"]] = blob
        else:
            for source in RAW_SOURCES:
                if args.source_cache:
                    blob = safe_local(args.source_cache.expanduser().resolve(), source["filename"]).read_bytes()
                    verify_bytes(blob, source, source["filename"])
                    network_meta = {"network_used": False}
                else:
                    blob, network_meta = fetch_source(source)
                write_new(output / "local_only/raw" / source["filename"], blob)
                receipt["acquisitions"].append(dict(source, **network_meta))
                if source["name"] != "SHANGHAI_D4":
                    data, metadata = reported_development(blob)
                    if (metadata["source_records"] != source["records"]
                            or metadata["source_start"] != source["start"] or metadata["source_end"] != source["end"]):
                        raise ValueError("Reported source metadata mismatch")
                    derived["development/" + source["name"] + ".csv"] = data
                    receipt["source_metadata"].append(dict(name=source["name"], **metadata))
                else:
                    workbook = extract_shanghai(blob)
                    write_new(output / "local_only/raw/mendeley_d4_carbon.xlsx", workbook)
                    derived.update(shanghai_partitions(workbook))
        for spec in DATA_FILES:
            blob = derived[spec["path"]]
            verify_bytes(blob, spec, spec["path"])
            write_new(output / spec["path"], blob)
            receipt["files"].append(spec)
        receipt.update(status="PASS", development_records=5864, independent_confirmation_records=261,
                       previously_available_2024_suffix_records_not_used=178,
                       development_end_exclusive="2024-01-01", confirmation_end_exclusive="2026-01-01",
                       figshare_doi="10.6084/m9.figshare.28682348.v1", shanghai_doi="10.17632/d4bv52fgmd.1",
                       licenses="Both deposits declare CC BY 4.0; attribution and source limitations in DATA_METHODS_ZH.md")
    except Exception as error:
        receipt.update(status="FAILED_STOPPED", error_type=type(error).__name__, error=str(error))
        raise
    finally:
        receipt["finished_utc"] = dt.datetime.now(dt.timezone.utc).isoformat()
        write_new(output / "provenance/DATA_INSTALL_MANIFEST.json", json_bytes(receipt))
    print(json.dumps(dict(status="PASS", output=str(output), development_records=5864,
                          confirmation_records=261, price_values_printed=False), ensure_ascii=False))


if __name__ == "__main__":
    main()
